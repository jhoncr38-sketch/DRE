# -*- coding: utf-8 -*-
"""
Recalcula as fórmulas do xlsx e exporta o PDF usando LibreOffice.

Uso:  python3 src/exportar.py            # usa build/final.xlsx
      python3 src/exportar.py caminho.xlsx

Requer LibreOffice instalado (`soffice` no PATH).
Gráficos, imagens e proteção sobrevivem ao recálculo.
"""
import pathlib
import shutil
import subprocess
import sys
import tempfile

RAIZ = pathlib.Path(__file__).resolve().parent.parent
ENTREGA = RAIZ / "entregaveis"


def soffice() -> str:
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        sys.exit("LibreOffice não encontrado. Instale-o ou ajuste o PATH.")
    return exe


def converter(entrada: pathlib.Path, formato: str, destino: pathlib.Path) -> pathlib.Path:
    """Converte via LibreOffice usando um perfil temporário (evita conflito de instância)."""
    destino.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as perfil:
        subprocess.run(
            [soffice(), "--headless", f"-env:UserInstallation=file://{perfil}",
             "--convert-to", formato, "--outdir", str(destino), str(entrada)],
            check=True, capture_output=True, timeout=300,
        )
    return destino / f"{entrada.stem}.{formato.split(':')[0]}"


def main() -> None:
    entrada = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / "build" / "final.xlsx"
    if not entrada.exists():
        sys.exit(f"Arquivo não encontrado: {entrada}\nRode antes: build.py, add_charts.py, finalize.py")

    # 1) recalcula: converte para xlsx num diretório temporário e traz de volta
    with tempfile.TemporaryDirectory() as tmp:
        recalculado = converter(entrada, "xlsx", pathlib.Path(tmp))
        final_xlsx = ENTREGA / "DRE_Aragao_Apresentacao.xlsx"
        ENTREGA.mkdir(exist_ok=True)
        shutil.copy(recalculado, final_xlsx)
    print(f"recalculado: {final_xlsx}")

    # 2) exporta o PDF a partir do arquivo já recalculado
    pdf = converter(final_xlsx, "pdf", ENTREGA)
    print(f"pdf: {pdf}")

    # 3) confere se sobrou algum erro de fórmula
    try:
        import openpyxl
    except ImportError:
        return
    wb = openpyxl.load_workbook(final_xlsx, data_only=True)
    erros = [
        f"{ws.title}!{c.coordinate}={c.value}"
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for c in row
        if isinstance(c.value, str) and c.value.startswith("#")
    ]
    print(f"erros de fórmula: {len(erros)}")
    for e in erros[:10]:
        print("  ", e)


if __name__ == "__main__":
    main()
