# -*- coding: utf-8 -*-
import pathlib as _pl
_RAIZ = _pl.Path(__file__).resolve().parent.parent
_ASSETS = _RAIZ / "assets"
_OUT = _RAIZ / "build"
_OUT.mkdir(exist_ok=True)

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

# ---------- Brand palette ----------
RED      = "A80201"   # brand red
CHAR     = "262626"   # charcoal
DARK     = "3F3F3F"
LGRAY    = "F4F4F4"   # very light band
MGRAY    = "E7E7E7"   # section band
REDTINT  = "F6E5E5"   # light red tint
YELLOW   = "FFF3CC"   # input cell
WHITE    = "FFFFFF"
LINE     = "D0D0D0"

FONT = "Arial"

def F(sz=10, b=False, color=CHAR, italic=False):
    return Font(name=FONT, size=sz, bold=b, color=color, italic=italic)

def fill(hexc):
    return PatternFill("solid", fgColor=hexc)

thin = Side(style="thin", color=LINE)
med_red = Side(style="medium", color=RED)
med_char = Side(style="medium", color=CHAR)

CUR = '[$-416]"R$" #,##0.00'  # currency (pt-BR)
PCT = '[$-416]0.0%'          # fraction -> percent (pt-BR)
PCTN = '[$-416]0.00"%"'     # percent points (pt-BR)
PCTF3 = '[$-416]0.000%'

wb = openpyxl.Workbook()

# =========================================================================
#  SHEET 0 : CONFIGURAÇÃO (alimenta os cabeçalhos das demais abas)
# =========================================================================
cfg = wb.active
cfg.title = "CONFIGURAÇÃO"
cfg.sheet_view.showGridLines = False
cfg.column_dimensions['A'].width = 30
cfg.column_dimensions['B'].width = 46
cfg.merge_cells('A1:B1')
t = cfg['A1']; t.value = "CONFIGURAÇÃO DO RELATÓRIO"
t.fill = PatternFill("solid", fgColor="A80201")
t.font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
t.alignment = Alignment(horizontal='left', vertical='center', indent=1)
cfg.row_dimensions[1].height = 26
_cfg_rows = [
    ("Nome da empresa", "ARAGÃO PRODUTOS FARMACÊUTICOS"),
    ("Período anual", "Exercício Anual"),
    ("Mês de referência", "Dezembro"),
]
for i,(k,v) in enumerate(_cfg_rows, start=3):
    a=cfg.cell(row=i,column=1,value=k); b=cfg.cell(row=i,column=2,value=v)
    a.font=Font(name="Arial",size=10,bold=True,color="262626")
    a.alignment=Alignment(horizontal='left',vertical='center',indent=1)
    b.font=Font(name="Arial",size=10,color="262626")
    b.alignment=Alignment(horizontal='left',vertical='center',indent=1)
    b.fill=PatternFill("solid",fgColor="FFF3CC")
    cfg.row_dimensions[i].height=19
h=cfg.cell(row=7,column=1,value="Altere os campos amarelos: os cabeçalhos das abas mudam sozinhos.")
h.font=Font(name="Arial",size=9,italic=True,color="6C6F73")
cfg.sheet_properties.tabColor="262626"
CFG_EMP="CONFIGURAÇÃO!$B$3"; CFG_ANO="CONFIGURAÇÃO!$B$4"; CFG_MES="CONFIGURAÇÃO!$B$5"

# =========================================================================
#  SHEET 1 : ANUAL
# =========================================================================
ws = wb.create_sheet("ANUAL")
ws.sheet_view.showGridLines = False

# column widths
ws.column_dimensions['A'].width = 44
ws.column_dimensions['B'].width = 17
ws.column_dimensions['C'].width = 12
for col in ['D','E','F']:
    ws.column_dimensions[col].width = 3
ws.column_dimensions['H'].width = 30
ws.column_dimensions['I'].width = 15

# ---- Title band (white with red rule) rows 1-4 ----
for r in range(1,5):
    ws.row_dimensions[r].height = 19
ws.merge_cells('A1:F4')
band = ws['A1']
band.fill = fill(WHITE)
band.alignment = Alignment(horizontal='right', vertical='center', wrap_text=True)
band.value = f'="DEMONSTRAÇÃO DE RESULTADO — DRE"&CHAR(10)&{CFG_EMP}&CHAR(10)&{CFG_ANO}' 
band.font = F(16, b=True, color=CHAR)
# red rule under band (row 5 thin)
ws.row_dimensions[5].height = 6
for c in range(1,7):
    cell = ws.cell(row=5, column=c)
    cell.fill = fill(RED)

# logo
img = XLImage(str(_ASSETS / "logo_crop.png"))
img.height = 74
img.width  = int(74 * 228/260)
ws.add_image(img, "A1")

# ---- column headers row 6 ----
hdr_row = 6
ws.row_dimensions[hdr_row].height = 20
headers = [("A","DESCRIÇÃO","left"),("B","VALOR (R$)","right"),("C","% RECEITA","right")]
for col,txt,al in headers:
    c = ws[f"{col}{hdr_row}"]
    c.value = txt
    c.fill = fill(RED)
    c.font = F(10, b=True, color=WHITE)
    c.alignment = Alignment(horizontal=al, vertical='center')

# helper to write a DRE line
key = {}
def line(r, label, value=None, role="item", numfmt=CUR, pct=None, indent=1, note=None):
    a = ws.cell(row=r, column=1, value=label)
    b = ws.cell(row=r, column=2)
    c = ws.cell(row=r, column=3)
    a.alignment = Alignment(horizontal='left', vertical='center', indent=indent)
    b.alignment = Alignment(horizontal='right', vertical='center')
    c.alignment = Alignment(horizontal='right', vertical='center')
    b.number_format = numfmt
    c.number_format = PCT
    if value is not None:
        b.value = value
    if pct is not None:
        c.value = pct
    # role styling
    if role == "section":
        a.font = F(10, b=True, color=WHITE); a.fill = fill(CHAR)
        b.fill = fill(CHAR); c.fill = fill(CHAR)
        b.font = F(10, b=True, color=WHITE); c.font = F(10, b=True, color=WHITE)
        a.alignment = Alignment(horizontal='left', vertical='center', indent=0)
    elif role == "subtotal":
        a.font = F(10, b=True, color=CHAR); a.fill = fill(MGRAY)
        b.fill = fill(MGRAY); c.fill = fill(MGRAY)
        b.font = F(10, b=True, color=CHAR); c.font = F(10, b=True, color=CHAR)
        a.alignment = Alignment(horizontal='left', vertical='center', indent=0)
        for cc in (a,b,c):
            cc.border = Border(top=thin, bottom=thin)
    elif role == "revenue":
        a.font = F(11, b=True, color=WHITE); a.fill = fill(RED)
        b.fill = fill(RED); c.fill = fill(RED)
        b.font = F(11, b=True, color=WHITE); c.font = F(11, b=True, color=WHITE)
        a.alignment = Alignment(horizontal='left', vertical='center', indent=0)
    elif role == "gross":
        a.font = F(10, b=True, color=RED); a.fill = fill(REDTINT)
        b.fill = fill(REDTINT); c.fill = fill(REDTINT)
        b.font = F(10, b=True, color=RED); c.font = F(10, b=True, color=RED)
        a.alignment = Alignment(horizontal='left', vertical='center', indent=0)
        for cc in (a,b,c):
            cc.border = Border(top=thin, bottom=thin)
    elif role == "result":
        a.font = F(12, b=True, color=WHITE); a.fill = fill(CHAR)
        b.fill = fill(CHAR); c.fill = fill(CHAR)
        b.font = F(12, b=True, color=WHITE); c.font = F(12, b=True, color=WHITE)
        a.alignment = Alignment(horizontal='left', vertical='center', indent=0)
        for cc in (a,b,c):
            cc.border = Border(top=med_red, bottom=med_red)
    else:  # item
        a.font = F(10, color=CHAR); b.font = F(10, color=CHAR); c.font = F(10, color=DARK)
    ws.row_dimensions[r].height = 15
    if note:
        a.comment = Comment(note, "Alexandre Araújo")
    return r

# ---- Build ANUAL body ----
r = 7
REC = r
line(r, "Receita total declarada (PGDAS)", 1187289.06, role="revenue"); ws[f"C{r}"].value = "=B{0}/$B${0}".format(REC)
r += 1
CUSTOS = r
line(r, "( - ) Custos", role="subtotal")
ws[f"B{r}"].value = f"=SUM(B{r+1}:B{r+2})"
ws[f"C{r}"].value = f"=B{r}/$B${REC}"
r += 1
line(r, "Custo de medicamentos, produtos e outros", 696475.30)
r += 1
line(r, "Produtos vencidos", 3160.96); ws[f"C{r}"].value = f"=B{r}/$B${REC}"
r += 1
GROSS = r
line(r, "(=) Lucro Operacional Bruto", role="gross")
ws[f"B{r}"].value = f"=B{REC}-B{CUSTOS}"
ws[f"C{r}"].value = f"=B{r}/$B${REC}"
r += 1
line(r, "( - ) Despesas Operacionais", role="section")
r += 1
# Despesas Gerais e Administrativas
GA = r
line(r, "Despesas Gerais e Administrativas", role="subtotal")
ga_first = r+1
ga_items = [
    ("Retirada de Sócio / Plano", 19031.85),
    ("Funcionários", 160841.13),
    ("Alimentação", None),
    ("Aluguel", 32400.0),
    ("Seguro", 1385.85),
    ("Telefone / Internet", 815.0),
    ("Energia", 20707.37),
    ("Cagece", 3396.07),
    ("Contabilidade", 15300.0),
    ("Corpvs Segurança", 1809.03),
    ("Taxa de banco e cartões", 10171.97),
    ("Sistema", 17790.53),
    ("Manutenção, Dedetização, Equipamentos", 240.0),
    ("Fatura Cartão", 23592.58),
    ("Outras despesas", 691.50),
]
r += 1
for lbl,val in ga_items:
    line(r, lbl, val)
    if val is not None:
        ws[f"C{r}"].value = f"=B{r}/$B${REC}"
    r += 1
ga_last = r-1
ws[f"B{GA}"].value = f"=SUM(B{ga_first}:B{ga_last})"
ws[f"C{GA}"].value = f"=B{GA}/$B${REC}"
# Despesas com Vendas
VEN = r
line(r, "Despesas com Vendas", role="subtotal")
ven_first = r+1
ven_items = [
    ("Publicidade e Propaganda", None),
    ("Combustível", None),
    ("Estacionamento / Lavagem", None),
    ("Seguro do carro", None),
    ("Acessórios / conserto veículo / revisão", 96.0),
]
r += 1
for lbl,val in ven_items:
    line(r, lbl, val)
    if val is not None:
        ws[f"C{r}"].value = f"=B{r}/$B${REC}"
    r += 1
ven_last = r-1
ws[f"B{VEN}"].value = f"=SUM(B{ven_first}:B{ven_last})"
ws[f"C{VEN}"].value = f"=B{VEN}/$B${REC}"
# Despesas Tributárias
TRIB = r
line(r, "Despesas Tributárias", role="subtotal")
trib_first = r+1
trib_items = [
    ("Simples Nacional", 60625.84),
    ("INSS", 8249.04),
    ("FGTS", 16118.12),
    ("SEFAZ", 2869.60),
    ("Taxas e Impostos Diversos - Cartório Aguiar", 14125.77),
    ("Parcelamento Simples", 1126.20),
    ("IPTU", 773.87),
    ("Conselho Regional de Farmácia", 2051.69),
    ("Novo Parcelamento Simples e-CAC", 1539.39),
    ("Novo Parcelamento Previdência", 511.76),
    ("Novo Parcelamento PGFN", 1025.34),
]
r += 1
for lbl,val in trib_items:
    line(r, lbl, val)
    if val is not None:
        ws[f"C{r}"].value = f"=B{r}/$B${REC}"
    r += 1
trib_last = r-1
ws[f"B{TRIB}"].value = f"=SUM(B{trib_first}:B{trib_last})"
ws[f"C{TRIB}"].value = f"=B{TRIB}/$B${REC}"
# spacer
ws.row_dimensions[r].height = 6
r += 1
# Resultado Líquido (value informed by client - keep as-is, flag discrepancy)
RES = r
line(r, "(=) RESULTADO LÍQUIDO", 81289.32, role="result",
     note=("Valor informado na apuração. A soma das linhas do DRE resulta em "
           "R$ 70.367,30 (diferença de R$ 10.922,02). Mantido o valor informado — "
           "verificar se há lançamento não detalhado."))
ws[f"C{r}"].value = f"=B{r}/$B${REC}"
r += 2
# Estoque section
line(r, "ESTOQUE / COMPRAS", role="section"); r += 1
COMPRA = r
line(r, "Compra de mercadoria (total ano)", 721324.99)
ws[f"C{r}"].value = f"=B{r}/$B${REC}"
last_body = r

# ---- helper table for charts (cols H/I, outside print area) ----
ws["H2"] = "AUX GRÁFICOS (não imprime)"; ws["H2"].font = F(8, italic=True, color=DARK)
pie_cats = [("Custos", CUSTOS), ("Despesas G&A", GA), ("Despesas c/ Vendas", VEN),
            ("Despesas Tributárias", TRIB), ("Resultado Líquido", RES)]
pr = 4
pie_first = pr
for lbl,src in pie_cats:
    ws.cell(row=pr, column=8, value=lbl).font = F(9)
    ws.cell(row=pr, column=9, value=f"=B{src}").number_format = CUR
    ws.cell(row=pr, column=9).font = F(9)
    pr += 1
pie_last = pr-1
pr += 1
col_first = pr
col_rows = [("Receita", f"=B{REC}"), ("(-) Custos", f"=B{CUSTOS}"),
            ("Lucro Op. Bruto", f"=B{GROSS}"),
            ("(-) Despesas", f"=B{GA}+B{VEN}+B{TRIB}"),
            ("Resultado Líquido", f"=B{RES}")]
for lbl,frm in col_rows:
    ws.cell(row=pr, column=8, value=lbl).font = F(9)
    ws.cell(row=pr, column=9, value=frm).number_format = CUR
    ws.cell(row=pr, column=9).font = F(9)
    pr += 1
col_last = pr-1

anual_meta = dict(REC=REC, RES=RES, last_body=last_body,
                  pie_first=pie_first, pie_last=pie_last,
                  col_first=col_first, col_last=col_last)

# =========================================================================
#  SHEET 2 : MÊS-DEZEMBRO
# =========================================================================
ws2 = wb.create_sheet("MÊS-DEZEMBRO")
ws2.sheet_view.showGridLines = False
ws2.column_dimensions['A'].width = 40
ws2.column_dimensions['B'].width = 14
ws2.column_dimensions['C'].width = 11
ws2.column_dimensions['D'].width = 10
ws2.column_dimensions['E'].width = 13
ws2.column_dimensions['F'].width = 3
ws2.column_dimensions['G'].width = 26
ws2.column_dimensions['H'].width = 13
ws2.column_dimensions['I'].width = 10
ws2.column_dimensions['J'].width = 13
ws2.column_dimensions['N'].width = 30
ws2.column_dimensions['O'].width = 14

for rr in range(1,5):
    ws2.row_dimensions[rr].height = 21
ws2.merge_cells('A1:J4')
b2 = ws2['A1']
b2.fill = fill(WHITE)
b2.alignment = Alignment(horizontal='right', vertical='center', wrap_text=True)
b2.value = f'="DEMONSTRAÇÃO DE RESULTADO — DRE"&CHAR(10)&{CFG_EMP}&CHAR(10)&"Mês de referência: "&{CFG_MES}' 
b2.font = F(16, b=True, color=CHAR)
ws2.row_dimensions[5].height = 6
for c in range(1,11):
    ws2.cell(row=5, column=c).fill = fill(RED)
img2 = XLImage(str(_ASSETS / "logo_crop.png"))
img2.height = 74
img2.width  = int(74 * 228/260)
ws2.add_image(img2, "A1")

# column headers row6 (left DRE)
ws2.row_dimensions[6].height = 20
for col,txt,al in [("A","DESCRIÇÃO","left"),("B","VALOR (R$)","right"),
                   ("C","% C/ CRÉD.","right"),("D","ALÍQUOTA","right"),("E","CRÉDITO (R$)","right")]:
    c = ws2[f"{col}6"]; c.value = txt; c.fill = fill(RED); c.font = F(10,b=True,color=WHITE)
    c.alignment = Alignment(horizontal=al, vertical='center')

def line2(r, label, value=None, credito=None, role="item", indent=1, note=None,
          bnumfmt=CUR, cnumfmt=CUR):
    a = ws2.cell(row=r, column=1, value=label)
    b = ws2.cell(row=r, column=2)
    c = ws2.cell(row=r, column=5)      # crédito em R$ agora na coluna E
    extra = [ws2.cell(row=r, column=3), ws2.cell(row=r, column=4)]
    a.alignment = Alignment(horizontal='left', vertical='center', indent=indent)
    b.alignment = Alignment(horizontal='right', vertical='center')
    c.alignment = Alignment(horizontal='right', vertical='center')
    for x in extra: x.alignment = Alignment(horizontal='right', vertical='center')
    b.number_format = bnumfmt; c.number_format = cnumfmt
    if value is not None: b.value = value
    if credito is not None: c.value = credito
    if role == "section":
        a.font=F(10,b=True,color=WHITE); a.fill=fill(CHAR); b.fill=fill(CHAR); c.fill=fill(CHAR)
        b.font=F(10,b=True,color=WHITE); c.font=F(10,b=True,color=WHITE)
        for x in extra: x.fill=fill(CHAR)
        a.alignment=Alignment(horizontal='left',vertical='center',indent=0)
    elif role=="subtotal":
        a.font=F(10,b=True,color=CHAR); a.fill=fill(MGRAY); b.fill=fill(MGRAY); c.fill=fill(MGRAY)
        b.font=F(10,b=True,color=CHAR); c.font=F(10,b=True,color=CHAR)
        for x in extra: x.fill=fill(MGRAY)
        for cc in (a,b,c)+tuple(extra): cc.border=Border(top=thin,bottom=thin)
    elif role=="revenue":
        a.font=F(11,b=True,color=WHITE); a.fill=fill(RED); b.fill=fill(RED); c.fill=fill(RED)
        b.font=F(11,b=True,color=WHITE); c.font=F(11,b=True,color=WHITE)
        for x in extra: x.fill=fill(RED)
        a.alignment=Alignment(horizontal='left',vertical='center',indent=0)
    elif role=="gross":
        a.font=F(10,b=True,color=RED); a.fill=fill(REDTINT); b.fill=fill(REDTINT); c.fill=fill(REDTINT)
        b.font=F(10,b=True,color=RED); c.font=F(10,b=True,color=RED)
        for x in extra: x.fill=fill(REDTINT)
        a.alignment=Alignment(horizontal='left',vertical='center',indent=0)
        for cc in (a,b,c)+tuple(extra): cc.border=Border(top=thin,bottom=thin)
    elif role=="result":
        a.font=F(12,b=True,color=WHITE); a.fill=fill(CHAR); b.fill=fill(CHAR); c.fill=fill(CHAR)
        b.font=F(12,b=True,color=WHITE); c.font=F(12,b=True,color=WHITE)
        for x in extra: x.fill=fill(CHAR)
        a.alignment=Alignment(horizontal='left',vertical='center',indent=0)
        for cc in (a,b,c)+tuple(extra): cc.border=Border(top=med_red,bottom=med_red)
    else:
        a.font=F(10,color=CHAR); b.font=F(10,color=CHAR); c.font=F(10,color=DARK)
    ws2.row_dimensions[r].height=18
    if note: a.comment=Comment(note,"Alexandre Araújo")
    return r

r = 7
REC2 = r
line2(r, "Receita declarada (PGDAS)", 106419.89, role="revenue"); r+=1
CUST2 = r
line2(r, "( - ) Custos", role="subtotal")
ws2[f"B{r}"].value = f"=SUM(B{r+1}:B{r+3})"
r+=1
line2(r, "Custo de medicamentos", None); r+=1
line2(r, "Custo de vendas de produtos / outros", 58039.60); r+=1
line2(r, "Produtos vencidos", None); r+=1
GROSS2 = r
line2(r, "(=) Lucro Operacional Bruto", role="gross")
ws2[f"B{r}"].value = f"=B{REC2}-B{CUST2}"
r+=1
line2(r, "( - ) Despesas Operacionais", role="section"); r+=1
GA2 = r
line2(r, "Despesas Gerais e Administrativas", role="subtotal")
ga2_first = r+1
r+=1
# items with optional credito CBS formula (relative to that row's B)
ga2 = [
    # (label, valor, % da base com direito a crédito, alíquota CBS)
    ("Retirada de sócio", None, 0.0, 0.0),
    ("HapVida", 804.29, 0.0, 0.0),
    ("Folha", 12346.56, 0.0, 0.0),
    ("Aluguel", 2700.0, 0.30, 0.08),
    ("Seguro", 125.87, 1.00, 0.08),
    ("Telefone / Internet", 151.26, 1.00, 0.08),
    ("Energia / Cagece", 1946.13, 1.00, 0.08),
    ("Contabilidade", 2100.0, 0.70, 0.08),
    ("Corpvs Segurança", 152.84, 1.00, 0.08),
    ("Taxa de banco e cartões", 2537.14, 1.00, 0.08),
    ("Sistema", 800.0, 1.00, 0.08),
    ("Manutenção, Dedetização, Equipamentos", 5.0, 1.00, 0.08),
    ("Outras despesas", 181.0, 0.0, 0.0),
]
PCTB = '[$-416]0%'
PCTA = '[$-416]0.00%'
for lbl,val,pbase,aliq in ga2:
    line2(r, lbl, val)
    cB=ws2.cell(row=r,column=3,value=pbase); cA=ws2.cell(row=r,column=4,value=aliq)
    cB.number_format=PCTB; cA.number_format=PCTA
    cB.fill=fill(YELLOW); cA.fill=fill(YELLOW)
    cB.font=F(9,b=True,color=CHAR); cA.font=F(9,b=True,color=CHAR)
    ws2.cell(row=r,column=5).value = f"=B{r}*C{r}*D{r}"
    r+=1
ga2_last = r-1
ws2[f"B{GA2}"].value = f"=SUM(B{ga2_first}:B{ga2_last})"
ws2[f"E{GA2}"].value = f"=SUM(E{ga2_first}:E{ga2_last})"
VEN2 = r
line2(r, "Despesas com Vendas", role="subtotal")
# credito geral total = sum of credit column of G&A block

ven2_first = r+1
r+=1
for lbl in ["Publicidade e Propaganda","Combustível","Estacionamento / Lavagem","Seguro do carro","Acessórios / conserto veículo / revisão"]:
    line2(r, lbl, None); r+=1
ven2_last = r-1
ws2[f"B{VEN2}"].value = f"=SUM(B{ven2_first}:B{ven2_last})"
TRIB2 = r
line2(r, "Despesas Tributárias", role="subtotal")
trib2_first = r+1
r+=1
for lbl,val in [("Simples Nacional",4889.93),("INSS",986.67),("FGTS",987.19),("SEFAZ",678.16)]:
    line2(r, lbl, val); r+=1
trib2_last = r-1
ws2[f"B{TRIB2}"].value = f"=SUM(B{trib2_first}:B{trib2_last})"
SIMPLES2 = trib2_first  # Simples Nacional row
ws2.row_dimensions[r].height = 6; r+=1
RES2 = r
line2(r, "(=) RESULTADO LÍQUIDO", role="result")
ws2[f"B{r}"].value = f"=B{GROSS2}-B{GA2}-B{VEN2}-B{TRIB2}"
r+=2
# Estoque / compras section (left column)
line2(r, "ESTOQUE / COMPRAS", role="section"); r+=1
COMPRA2 = r
line2(r, "Compra de mercadoria (dezembro)", 62239.31); r+=1
r+=0
line2(r, "Inventário 25/11/2025 — Custo total", 211166.93); r+=1
line2(r, "Inventário 25/11/2025 — Preço total", 569370.33); r+=1
line2(r, "Estoque atualizado 26/11/2025 — Custo total", 177250.29); r+=1
line2(r, "Estoque atualizado 26/11/2025 — Preço total", 453222.22); r+=1
line2(r, "Pagamento a fornecedor até dezembro", 506100.89); r+=1
left_last = r-1

# ---------- Right-side box: Memória de cálculo CBS (reform) — cols G..J ----------
PCT2 = '[$-416]0.00%'
ws2.merge_cells('G6:J6')
hb = ws2['G6']; hb.value = "MEMÓRIA DE CÁLCULO — REFORMA (CBS)"
hb.fill=fill(CHAR); hb.font=F(10,b=True,color=WHITE)
hb.alignment=Alignment(horizontal='left',vertical='center')
for cc in ('H6','I6','J6'): ws2[cc].fill=fill(CHAR)

def _al(c,h='right'): c.alignment=Alignment(horizontal=h,vertical='center',indent=(1 if h=='left' else 0))

def rline(r, label, base=None, aliq=None, valor=None, hfmt=CUR, gfmt=None,
          base_edit=False, aliq_edit=False, h_edit=False, total=False, note=None):
    e=ws2.cell(row=r,column=7,value=label); fx=ws2.cell(row=r,column=8)
    gx=ws2.cell(row=r,column=9); hx=ws2.cell(row=r,column=10)
    e.font=F(9,color=CHAR); _al(e,'left')
    _al(fx); _al(gx); _al(hx)
    fx.number_format=CUR; gx.number_format=(gfmt or PCT2); hx.number_format=hfmt
    if base is not None: fx.value=base
    if aliq is not None: gx.value=aliq
    if valor is not None: hx.value=valor
    fx.font=F(9,color=DARK); gx.font=F(9,color=DARK); hx.font=F(9,color=DARK)
    if base_edit: fx.fill=fill(YELLOW); fx.font=F(9,b=True,color=CHAR)
    if aliq_edit: gx.fill=fill(YELLOW); gx.font=F(9,b=True,color=CHAR)
    if h_edit: hx.fill=fill(YELLOW); hx.font=F(9,b=True,color=CHAR)
    if total:
        e.font=F(9,b=True,color=CHAR); hx.font=F(9,b=True,color=CHAR)
        for c in (e,fx,gx,hx): c.border=Border(top=thin)
    ws2.row_dimensions[r].height=15
    if note: e.comment=Comment(note,"Alexandre Araújo")

rline(7,"Receita", valor=f"=B{REC2}")
rline(8,"Simples Nacional", valor=f"=B{SIMPLES2}"); F_SIMP=8
rline(9,"Alíquota efetiva (Simples ÷ Receita)", valor="=IF(J7=0,0,(J8/J7)*100)", hfmt=PCTN,
      note="Imposto efetivamente pago no mês dividido pela receita do mês.")
rline(10,"% da CBS sobre a efetiva", valor=0.1533, gfmt=PCT2, hfmt=PCT2, h_edit=True,
      note=("Percentual da alíquota efetiva correspondente à CBS. A memória RBT12 abaixo usa 15,5% — "
            "alinhar os dois critérios."))
rline(11,"Percentual da CBS", valor="=J9*J10", hfmt=PCTN)
rline(12,"Percentual ICMS", valor="=J9*33.5%", hfmt=PCTN)
rline(13,"Alíquota efetiva (líq. ICMS)", valor="=J9-J12", hfmt=PCTN)
rline(14,"Valor de CBS", valor="=J7*(J11/100)")
rline(15,"Simples sem a CBS", valor=f"=J8-J14",
      note=("Simples Nacional do mês menos a parcela correspondente à CBS. "
            "É o valor que permaneceria no Simples caso a CBS passe a ser recolhida à parte."))
ws2.cell(row=15,column=7).font=F(9,b=True,color=CHAR)
ws2.cell(row=15,column=10).font=F(9,b=True,color=CHAR)
for _c in (7,8,9,10): ws2.cell(row=15,column=_c).border=Border(top=thin)
F_SIMPSEMCBS=15

for col,txt in [(8,"Base (R$)"),(9,"Alíquota"),(10,"Valor (R$)")]:
    c=ws2.cell(row=16,column=col,value=txt); c.font=F(8,b=True,color=DARK); _al(c)
sh=ws2.cell(row=16,column=7,value="DÉBITO — receita tributada"); sh.font=F(8,b=True,color=DARK); _al(sh,'left')

# 4 linhas de categoria de receita (nome/base/alíquota editáveis)
DEB_FIRST, DEB_LAST = 17, 20
_deb=[("Medicamentos",103685.52,0.032),("Demais itens",2734.37,0.08),("",None,None),("",None,None)]
for k,(nm,bs,al) in enumerate(_deb):
    r=DEB_FIRST+k
    rline(r, nm, base=bs, aliq=al, valor=f"=IF(OR(H{r}=\"\",I{r}=\"\"),0,H{r}*I{r})",
          base_edit=True, aliq_edit=True)
    e=ws2.cell(row=r,column=7); e.fill=fill(YELLOW); e.font=F(9,b=True,color=CHAR)
rline(21,"Débito total", valor=f"=SUM(J{DEB_FIRST}:J{DEB_LAST})", total=True); F_DTOT=21

gh=ws2.cell(row=23,column=7,value="CRÉDITO"); gh.font=F(8,b=True,color=DARK); _al(gh,'left')
rline(24,"Crédito geral (despesas)",
      base=f"=SUMPRODUCT(B{ga2_first}:B{ga2_last},C{ga2_first}:C{ga2_last})",
      aliq="=IF(H24=0,0,J24/H24)", valor=f"=E{GA2}",
      note="Detalhado item a item na DRE ao lado (colunas % c/ créd. e Alíquota).")
CRD_FIRST, CRD_LAST = 25, 28
_crd=[("Crédito medicamentos",f"=B{COMPRA2}*97%",0.032),("Crédito estoque",f"=B{COMPRA2}*3%",0.08),
      ("",None,None),("",None,None)]
for k,(nm,bs,al) in enumerate(_crd):
    r=CRD_FIRST+k
    rline(r, nm, base=bs, aliq=al, valor=f"=IF(OR(H{r}=\"\",I{r}=\"\"),0,H{r}*I{r})",
          base_edit=(bs is None), aliq_edit=True)
    e=ws2.cell(row=r,column=7); e.fill=fill(YELLOW); e.font=F(9,b=True,color=CHAR)
rline(29,"Saldo credor do mês anterior", valor=0, h_edit=True,
      note="Se o mês anterior fechou com crédito maior que o débito, lance aqui o saldo transportado.")
SALDO_ANT=29
rline(30,"Crédito total", valor=f"=J24+SUM(J{CRD_FIRST}:J{CRD_LAST})+J{SALDO_ANT}", total=True); F_CTOT=30

# saldo do mês: a pagar (>0) ou credor a transportar (<0)
e=ws2.cell(row=32,column=7,value="CBS A PAGAR"); hx=ws2.cell(row=32,column=10,value=f"=MAX(0,J{F_DTOT}-J{F_CTOT})")
e.fill=fill(REDTINT); hx.fill=fill(REDTINT)
ws2.cell(row=32,column=8).fill=fill(REDTINT); ws2.cell(row=32,column=9).fill=fill(REDTINT)
e.font=F(10,b=True,color=RED); hx.font=F(10,b=True,color=RED)
_al(e,'left'); _al(hx); hx.number_format=CUR
for col in (7,8,9,10): ws2.cell(row=32,column=col).border=Border(top=med_red,bottom=med_red)
F_CBSPAY=32
rline(33,"Saldo credor a transportar", valor=f"=MAX(0,J{F_CTOT}-J{F_DTOT})",
      note="Lance este valor como 'Saldo credor do mês anterior' na planilha do mês seguinte.")
ws2.cell(row=33,column=7).font=F(9,italic=True,color=DARK)
ws2.cell(row=33,column=10).font=F(9,italic=True,color=DARK)

ws2.merge_cells('G34:J34')
h=ws2.cell(row=34,column=7,value="CÁLCULO DA ALÍQUOTA EFETIVA (RBT12)")
h.fill=fill(MGRAY); h.font=F(9,b=True,color=CHAR); _al(h,'left')
for col in (8,9,10): ws2.cell(row=34,column=col).fill=fill(MGRAY)
rline(35,"RBT12", valor=1179636.95, h_edit=True)
rline(36,"Alíquota (nominal)", valor=0.107, hfmt=PCT, h_edit=True)
rline(37,"Parcela a deduzir (PD)", valor=22500.0, h_edit=True)
rline(38,"Alíquota efetiva", valor="=IF(J35=0,0,(((J35*J36)-J37)/J35)*100)", hfmt=PCTN)
rline(39,"CBS (parcela)", valor="=J38*15.5%", hfmt=PCTN)
rline(40,"ICMS (parcela)", valor="=J38*33.5%", hfmt=PCTN)
rline(42,"Valor avulso (planilha original)", valor=6287.09,
      note=("Valor 6.287,09 digitado na memória de cálculo original (célula F10), sem rótulo "
            "e sem ser usado por nenhuma fórmula. Preservado apenas para registro."))
ws2.cell(row=42,column=7).font=F(9,italic=True,color=DARK)
ws2.cell(row=42,column=10).font=F(9,italic=True,color=DARK)
right_last=42

# ---------- helper for charts (N/O, outside print) ----------
ws2["N2"]="AUX GRÁFICOS (não imprime)"; ws2["N2"].font=F(8,italic=True,color=DARK)
ws2["N4"]="IBS a pagar (informar)"; ws2["N4"].font=F(9)
ws2["O4"]=0; ws2["O4"].number_format=CUR; ws2["O4"].fill=fill(YELLOW); ws2["O4"].font=F(9,b=True)
IBS_CELL="O4"
ws2["N6"]="Hoje (Simples)"; ws2["N6"].font=F(9)
ws2["O6"]=f"=J{F_SIMP}"; ws2["O6"].number_format=CUR; ws2["O6"].font=F(9)
ws2["N7"]="Simples sem a CBS"; ws2["N7"].font=F(9)
ws2["O7"]=f"=J{F_SIMPSEMCBS}"; ws2["O7"].number_format=CUR; ws2["O7"].font=F(9)
ws2["N8"]="Na transição"; ws2["N8"].font=F(9)
ws2["O8"]=f"=J{F_SIMPSEMCBS}+J{F_CBSPAY}+{IBS_CELL}"; ws2["O8"].number_format=CUR; ws2["O8"].font=F(9)
cmp_first, cmp_last = 6,8
ws2["N10"]="Débito CBS"; ws2["N10"].font=F(9); ws2["O10"]=f"=J{F_DTOT}"; ws2["O10"].number_format=CUR; ws2["O10"].font=F(9)
ws2["N11"]="Crédito CBS"; ws2["N11"].font=F(9); ws2["O11"]=f"=J{F_CTOT}"; ws2["O11"].number_format=CUR; ws2["O11"].font=F(9)
ws2["N12"]="CBS a pagar"; ws2["N12"].font=F(9); ws2["O12"]=f"=J{F_CBSPAY}"; ws2["O12"].number_format=CUR; ws2["O12"].font=F(9)
mech_first, mech_last = 10,12

mes_meta = dict(REC2=REC2, RES2=RES2, left_last=left_last, right_last=right_last,
                cmp_first=cmp_first, cmp_last=cmp_last,
                mech_first=mech_first, mech_last=mech_last)

wb.save(str(_OUT / "styled_nochart.xlsx"))
import json
print(json.dumps({"anual":anual_meta,"mes":mes_meta}, ensure_ascii=False))
